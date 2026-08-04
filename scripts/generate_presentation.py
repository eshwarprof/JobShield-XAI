import sys
import os
import subprocess

# Ensure python-pptx is installed
try:
    import pptx
except ImportError:
    print("python-pptx not found. Installing via pip...")
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-pptx"])
    import pptx

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# -------------------------------------------------------------
# Premium Dark Color Palette Definitions
# -------------------------------------------------------------
COLOR_BG_DARK = RGBColor(15, 23, 42)        # Deep Slate 900 (Global Slide Background)
COLOR_CARD_BG = RGBColor(30, 41, 59)        # Slate 800 (Card background)
COLOR_CARD_BORDER = RGBColor(51, 65, 85)    # Slate 700 (Card border)
COLOR_TEXT_LIGHT = RGBColor(248, 250, 252)  # Slate 50 (Title / Body text)
COLOR_TEXT_MUTED = RGBColor(148, 163, 184)  # Slate 400 (Helper descriptions)

# Highlight Colors
COLOR_BLUE = RGBColor(59, 130, 246)         # Blue 500 (Primary accent / headers)
COLOR_TEAL = RGBColor(96, 165, 250)         # Blue 400 (Secondary accent / labels)
COLOR_GREEN = RGBColor(16, 185, 129)        # Emerald 500 (Legitimate status)
COLOR_RED = RGBColor(239, 68, 68)           # Red 500 (Fraud status / danger alerts)
COLOR_AMBER = RGBColor(245, 158, 11)        # Amber 500 (Warnings / highlighting)

# -------------------------------------------------------------
# Presentation Configuration
# -------------------------------------------------------------
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_slide_layout = prs.slide_layouts[6] 

def apply_theme(slide):
    """Sets a solid premium dark background on the slide."""
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = COLOR_BG_DARK

def add_header(slide, title_text):
    """Adds a standard premium header with an underline accent bar."""
    # Text container
    header_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12.133), Inches(0.8))
    tf = header_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.name = 'Trebuchet MS'
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_LIGHT

    # Underline accent bar
    accent_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 
        Inches(0.6), Inches(1.15), Inches(2.2), Inches(0.05)
    )
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = COLOR_BLUE
    accent_bar.line.color.rgb = COLOR_BLUE

def add_footer(slide, current_slide, total_slides=12):
    """Applies a standard footer with project name and slide count."""
    footer_box = slide.shapes.add_textbox(Inches(0.6), Inches(7.0), Inches(12.133), Inches(0.3))
    tf = footer_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = f"Explainable Fake Job Advertisement Detection Framework | Slide {current_slide} of {total_slides}"
    p.font.name = 'Calibri'
    p.font.size = Pt(10)
    p.font.color.rgb = COLOR_TEXT_MUTED
    p.alignment = PP_ALIGN.LEFT

def create_card(slide, left, top, width, height, title, content_list=[], border_color=COLOR_CARD_BORDER):
    """Utility helper to draw a modern glassmorphism-styled container card with bullet items."""
    # Background shape
    card_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card_shape.fill.solid()
    card_shape.fill.fore_color.rgb = COLOR_CARD_BG
    card_shape.line.color.rgb = border_color
    card_shape.line.width = Pt(1.5)

    tf = card_shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.25)
    tf.margin_top = Inches(0.2)
    tf.margin_bottom = Inches(0.2)

    # Title paragraph
    p_title = tf.paragraphs[0]
    p_title.text = title
    p_title.font.name = 'Trebuchet MS'
    p_title.font.size = Pt(18)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_TEAL
    p_title.space_after = Pt(12)

    # Content bullets
    for item in content_list:
        p_bullet = tf.add_paragraph()
        p_bullet.space_after = Pt(10)
        p_bullet.font.name = 'Calibri'
        p_bullet.font.size = Pt(13)
        p_bullet.font.color.rgb = COLOR_TEXT_LIGHT
        p_bullet.line_spacing = 1.15

        if isinstance(item, tuple):
            lead, desc = item
            run_l = p_bullet.add_run()
            run_l.text = f"•  {lead}: "
            run_l.font.bold = True
            run_l.font.color.rgb = COLOR_TEXT_LIGHT
            
            run_d = p_bullet.add_run()
            run_d.text = desc
            run_d.font.color.rgb = COLOR_TEXT_MUTED
        else:
            p_bullet.text = f"•  {item}"

    return card_shape

# =============================================================
# SLIDE 1: Title Slide
# =============================================================
slide_1 = prs.slides.add_slide(blank_slide_layout)
apply_theme(slide_1)

# Large Glowing Title
title_box = slide_1.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.733), Inches(2.2))
tf = title_box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Explainable Fake Job Advertisement\nDetection Framework"
p.font.name = 'Trebuchet MS'
p.font.size = Pt(44)
p.font.bold = True
p.font.color.rgb = COLOR_TEXT_LIGHT
p.alignment = PP_ALIGN.LEFT

p_sub = tf.add_paragraph()
p_sub.text = "A Trust-Aware, Auditable Decision Support System for Online Recruitment Fraud"
p_sub.font.name = 'Calibri'
p_sub.font.size = Pt(18)
p_sub.font.color.rgb = COLOR_BLUE
p_sub.space_before = Pt(12)

# Divider line
div = slide_1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(4.3), Inches(4.0), Inches(0.04))
div.fill.solid()
div.fill.fore_color.rgb = COLOR_BLUE
div.line.color.rgb = COLOR_BLUE

# Project stats/authorship
meta_box = slide_1.shapes.add_textbox(Inches(0.8), Inches(4.8), Inches(11.733), Inches(1.5))
tf_meta = meta_box.text_frame
tf_meta.word_wrap = True

meta_runs = [
    ("MODEL BASELINE", "Fraud-BERT (Discover Computing, 2025)"),
    ("CORE STACK", "React / Node.js / FastAPI / MongoDB"),
    ("INTEGRATION", "SHAP & LIME XAI, Heuristics Scorecard, OCR Ingestion")
]

p_m = tf_meta.paragraphs[0]
p_m.space_after = Pt(8)
for i, (label, val) in enumerate(meta_runs):
    if i > 0:
        p_m = tf_meta.add_paragraph()
        p_m.space_after = Pt(8)
    
    r_lbl = p_m.add_run()
    r_lbl.text = f"{label}: "
    r_lbl.font.name = 'Trebuchet MS'
    r_lbl.font.size = Pt(11)
    r_lbl.font.bold = True
    r_lbl.font.color.rgb = COLOR_TEAL
    
    r_val = p_m.add_run()
    r_val.text = val
    r_val.font.name = 'Calibri'
    r_val.font.size = Pt(12)
    r_val.font.color.rgb = COLOR_TEXT_MUTED

add_footer(slide_1, 1)

# =============================================================
# SLIDE 2: Problem Statement & Motivation
# =============================================================
slide_2 = prs.slides.add_slide(blank_slide_layout)
apply_theme(slide_2)
add_header(slide_2, "1. Introduction & Motivation")

# Left Column Card (Motivation)
create_card(
    slide_2, Inches(0.6), Inches(1.6), Inches(5.8), Inches(5.0),
    "Problem Statement",
    [
        ("Explosion of Recruitment Fraud", "The rapid growth of online job hunting has made job seekers vulnerable to sophisticated spoofing schemes and fake postings."),
        ("Significant Candidate Harm", "Leads to severe outcomes including financial extortion (fees) and harvesting of personal identification info."),
        ("Opaque Black-Box Models", "Modern classifiers predict fraud probabilities but offer zero explanation of results, leaving candidates unable to verify details.")
    ]
)

# Right Column Card (Our Approach)
create_card(
    slide_2, Inches(6.8), Inches(1.6), Inches(5.9), Inches(5.0),
    "Proposed Project Vision",
    [
        ("Actionable Interpretability", "Move from opaque binary classifications to detailed warnings showing exactly why a post is flagged as suspicious."),
        ("Domain-Specific Heuristics", "Cross-validate model evaluations against rule-based indicators (e.g. registration fee requests, free domains)."),
        ("Auditable Trust Decisions", "Empower candidate decision support with dynamic explanation panels (SHAP / LIME) and risk badges."),
        ("Ingestion Adaptability", "Provide uniform scanning pipelines for structured URLs, raw text, documents (PDF), and screenshot images (OCR).")
    ],
    border_color=COLOR_BLUE
)

add_footer(slide_2, 2)

# =============================================================
# SLIDE 3: Baseline vs. Proposed Extension
# =============================================================
slide_3 = prs.slides.add_slide(blank_slide_layout)
apply_theme(slide_3)
add_header(slide_3, "2. Base Paper vs. Proposed Framework")

rows, cols = 7, 3
left, top, width, height = Inches(0.6), Inches(1.6), Inches(12.133), Inches(5.0)
table_shape = slide_3.shapes.add_table(rows, cols, left, top, width, height)
table = table_shape.table

# Align Column Widths
table.columns[0].width = Inches(2.2)
table.columns[1].width = Inches(4.8)
table.columns[2].width = Inches(5.133)

headers = ["Feature", "Fraud-BERT Baseline (Taneja et al., 2025)", "Proposed Extension Framework"]
for idx, text in enumerate(headers):
    cell = table.cell(0, idx)
    cell.text = text
    cell.fill.solid()
    cell.fill.fore_color.rgb = COLOR_CARD_BG
    p = cell.text_frame.paragraphs[0]
    p.font.name = 'Trebuchet MS'
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEAL
    p.alignment = PP_ALIGN.CENTER

comparison_rows = [
    ("ML Model Core", "Standard BERT-base-uncased classifier.", "BERT-base-uncased with tuned decision boundaries."),
    ("Interpretability", "None (Black Box). Labels only.", "Interactive SHAP (Global population features) and LIME (Local word explanations)."),
    ("Trust Indicators", "None.", "0-100 deterministic Trust Score based on 8 heuristics (domain, logo, contact checks)."),
    ("Severity Index", "None.", "0-100 danger indicator detecting money requests and WhatsApp/Telegram recruitments."),
    ("Risk Tiering", "None.", "4-tier Risk Badge (Low, Medium, High, Critical) mapping BERT prob and Trust Score."),
    ("Ingestion Modes", "Text fields only.", "Supports URL scraping, PDF extraction, screenshot image OCR, and raw text.")
]

for r_idx, row in enumerate(comparison_rows, start=1):
    for c_idx, val in enumerate(row):
        cell = table.cell(r_idx, c_idx)
        cell.text = val
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_CARD_BG if r_idx % 2 == 0 else COLOR_BG_DARK
        p = cell.text_frame.paragraphs[0]
        p.font.name = 'Calibri'
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_TEXT_LIGHT
        if c_idx == 0:
            p.font.bold = True
            p.font.color.rgb = COLOR_TEAL

add_footer(slide_3, 3)

# =============================================================
# SLIDE 4: System Architecture
# =============================================================
slide_4 = prs.slides.add_slide(blank_slide_layout)
apply_theme(slide_4)
add_header(slide_4, "3. System Architecture Design")

# Card 1: React UI
create_card(
    slide_4, Inches(0.6), Inches(1.6), Inches(3.8), Inches(3.4),
    "Frontend (React 19)",
    [
        "Ingests job links, raw text, PDF files, and image files.",
        "Displays the interactive Explainability Dashboard.",
        "Renders LIME highlighting and SHAP feature panels.",
        "Connects only to Node.js backend (strict separation)."
    ]
)

# Card 2: Express gateway
create_card(
    slide_4, Inches(4.7), Inches(1.6), Inches(3.9), Inches(3.4),
    "Backend (Node/Express)",
    [
        "Authentication & session management.",
        "Enforces rate limiting and input validation rules.",
        "Logs query details and transaction results to MongoDB.",
        "Proxies verified requests to Python FastAPI AI service."
    ],
    border_color=COLOR_BLUE
)

# Card 3: Python AI
create_card(
    slide_4, Inches(8.9), Inches(1.6), Inches(3.8), Inches(3.4),
    "AI Service (FastAPI)",
    [
        "Executes tokenization and clean preprocessing.",
        "Fine-tuned BERT classification inference.",
        "Calculates Trust, Severity and Risk Score metrics.",
        "Precomputes SHAP cache; runs LIME iterations on request."
    ]
)

# Bottom Detail Bar
bottom_shape = slide_4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(5.3), Inches(12.133), Inches(1.4))
bottom_shape.fill.solid()
bottom_shape.fill.fore_color.rgb = COLOR_CARD_BG
bottom_shape.line.color.rgb = COLOR_CARD_BORDER
bottom_shape.line.width = Pt(1.5)

tf = bottom_shape.text_frame
tf.word_wrap = True
tf.margin_left = tf.margin_right = Inches(0.2)
tf.margin_top = tf.margin_bottom = Inches(0.15)
p_b = tf.paragraphs[0]
p_b.text = "Decoupled Data Architecture & Deployment Setup"
p_b.font.name = 'Trebuchet MS'
p_b.font.size = Pt(15)
p_b.font.bold = True
p_b.font.color.rgb = COLOR_TEAL
p_b.space_after = Pt(4)

p_b_body = tf.add_paragraph()
p_b_body.text = "• Database Layer: MongoDB houses 'users' (auth details) and 'analyses' schemas (saving predictions and audit records).\n• Performance Caching: The BERT model classifier and the precomputed global SHAP values are loaded once on FastAPI startup to avoid local execution bottlenecks."
p_b_body.font.name = 'Calibri'
p_b_body.font.size = Pt(12)
p_b_body.font.color.rgb = COLOR_TEXT_LIGHT
p_b_body.line_spacing = 1.15

add_footer(slide_4, 4)

# =============================================================
# SLIDE 5: Dataset & Preprocessing (Light Details)
# =============================================================
slide_5 = prs.slides.add_slide(blank_slide_layout)
apply_theme(slide_5)
add_header(slide_5, "4. Dataset & Data Imbalance Insights")

# Left Info Card
create_card(
    slide_5, Inches(0.6), Inches(1.6), Inches(5.8), Inches(5.0),
    "EMSCAD Dataset Properties",
    [
        ("Total Records", "17,880 postings mapped across 18 schema variables."),
        ("Severe Imbalance", "17,014 Legitimate postings vs. only 866 Fraudulent postings (4.84%)."),
        ("weighted loss", "To prevent classifier collapse (where models optimize simple accuracy by predicting only the majority class), we apply class weights in our cross-entropy loss function during BERT training."),
        ("Text Preprocessing", "Strips HTML tags, normalizes Unicode characters, aggregates nested fields (profile, requirements), and performs WordPiece tokenization.")
    ]
)

# Right: Visual table explaining missingness
table_box = slide_5.shapes.add_textbox(Inches(6.8), Inches(1.6), Inches(5.9), Inches(0.5))
tf_tbl = table_box.text_frame
p_t = tf_tbl.paragraphs[0]
p_t.text = "Missingness Percentages: Legit vs Fraud"
p_t.font.name = 'Trebuchet MS'
p_t.font.size = Pt(18)
p_t.font.bold = True
p_t.font.color.rgb = COLOR_TEAL

rows, cols = 5, 3
tbl_s = slide_5.shapes.add_table(rows, cols, Inches(6.8), Inches(2.2), Inches(5.933), Inches(2.6))
tbl = tbl_s.table
tbl.columns[0].width = Inches(2.333)
tbl.columns[1].width = Inches(1.8)
tbl.columns[2].width = Inches(1.8)

headers = ["Field Name", "Missing (Legit)", "Missing (Fraud)"]
for idx, text in enumerate(headers):
    cell = tbl.cell(0, idx)
    cell.text = text
    cell.fill.solid()
    cell.fill.fore_color.rgb = COLOR_CARD_BG
    p = cell.text_frame.paragraphs[0]
    p.font.name = 'Trebuchet MS'
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEAL
    p.alignment = PP_ALIGN.CENTER

missingness_rows = [
    ("company_profile", "15.99%", "67.78%"),
    ("salary_range", "84.45%", "74.25%"),
    ("benefits", "40.30%", "42.03%"),
    ("requirements", "14.95%", "17.78%")
]

for r_idx, row in enumerate(missingness_rows, start=1):
    for c_idx, val in enumerate(row):
        cell = tbl.cell(r_idx, c_idx)
        cell.text = val
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_CARD_BG if r_idx % 2 == 0 else COLOR_BG_DARK
        p = cell.text_frame.paragraphs[0]
        p.font.name = 'Calibri'
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_TEXT_LIGHT
        if c_idx == 0:
            p.font.bold = True
            p.font.color.rgb = COLOR_TEAL
            p.alignment = PP_ALIGN.LEFT
        else:
            p.alignment = PP_ALIGN.CENTER

# Right bottom note
note_box = slide_5.shapes.add_textbox(Inches(6.8), Inches(5.0), Inches(5.933), Inches(1.5))
tf_n = note_box.text_frame
tf_n.word_wrap = True
p_n = tf_n.paragraphs[0]
p_n.text = "Signal Validation Insight:"
p_n.font.name = 'Trebuchet MS'
p_n.font.size = Pt(14)
p_n.font.bold = True
p_n.font.color.rgb = COLOR_BLUE
p_n.space_after = Pt(4)

p_n_body = tf_n.add_paragraph()
p_n_body.text = "The significant difference in missing company profiles (67.78% in fraud vs. 15.99% in legitimate postings) confirms that lacking profile details is a strong warning sign. This directly validates the rule penalties used in the trust score scorecard."
p_n_body.font.name = 'Calibri'
p_n_body.font.size = Pt(12)
p_n_body.font.color.rgb = COLOR_TEXT_MUTED
p_n_body.line_spacing = 1.2

add_footer(slide_5, 5)

# =============================================================
# SLIDE 6: AI Classification Pipeline
# =============================================================
slide_6 = prs.slides.add_slide(blank_slide_layout)
apply_theme(slide_6)
add_header(slide_6, "5. AI Classification Pipeline Steps")

# 4 card process blocks
pipeline_steps = [
    ("1. Input Ingest", "Extracts text from scraped URLs, processes PDF files using PyMuPDF, or runs EasyOCR on screenshot images."),
    ("2. Preprocess", "Strips HTML characters, cleans emojis, normalizes whitespace formats, and generates WordPiece tokens."),
    ("3. ML Inference", "Fine-tuned BERT-base-uncased computes raw classification probability. Applies tuned 0.65 threshold."),
    ("4. Output Engines", "Runs heuristic calculations for Trust/Severity scorecards and runs live LIME explainability models.")
]

for idx, (title, text) in enumerate(pipeline_steps):
    left = Inches(0.6 + idx * 3.033)
    border_c = COLOR_BLUE if idx == 2 else COLOR_CARD_BORDER # Highlight ML phase
    create_card(
        slide_6, left, Inches(1.8), Inches(2.833), Inches(4.5),
        title,
        [text],
        border_color=border_c
    )

add_footer(slide_6, 6)

# =============================================================
# SLIDE 7: Explainable AI Engine
# =============================================================
slide_7 = prs.slides.add_slide(blank_slide_layout)
apply_theme(slide_7)
add_header(slide_7, "6. Explainable AI (XAI) Integration")

# Left Column Card: LIME
create_card(
    slide_7, Inches(0.6), Inches(1.6), Inches(5.8), Inches(5.0),
    "LIME (Local/Instance Explanations)",
    [
        ("Real-time Perturbation", "Runs live per evaluation request, perturbing input sentences locally to observe model classification score fluctuations."),
        ("Word-Level Highlighting", "Identifies specific words driving predictions, color-coding features for job seekers (Legitimate signals vs. Fraud signals)."),
        ("Transparency Focus", "Highlights suspicious terms like 'urgent', 'registration fee', and 'wire transfer' so applicants know exactly why a warning was raised.")
    ]
)

# Right Column Card: SHAP
create_card(
    slide_7, Inches(6.8), Inches(1.6), Inches(5.9), Inches(5.0),
    "SHAP (Global/Population Explanations)",
    [
        ("Population Analysis", "Computes cooperative game theory attribution scores across the entire EMSCAD training dataset."),
        ("Startup Optimization", "SHAP values are precomputed offline due to high overhead, cached on disk, and loaded immediately at application start."),
        ("System Verification", "Provides developers with an overview of what the BERT model learns overall, ensuring it avoids dataset bias or noise.")
    ],
    border_color=COLOR_BLUE
)

add_footer(slide_7, 7)

# =============================================================
# SLIDE 8: Trust Scoreboard
# =============================================================
slide_8 = prs.slides.add_slide(blank_slide_layout)
apply_theme(slide_8)
add_header(slide_8, "7. Deterministic Trust Score Engine")

# Left Column: Trust engine description
create_card(
    slide_8, Inches(0.6), Inches(1.6), Inches(5.8), Inches(5.0),
    "Transparent Trust Evaluation",
    [
        ("Anti-Stacking Approach", "Avoids placing a second opaque ML model on top of BERT. Instead, standard deterministic rules are used to score trust."),
        ("Heuristic Deductions", "Starts at 100 points. Penalties are deducted when specific anomalies are detected in text or metadata."),
        ("8 Active Indicators", "Checks contacts for free email domains (Gmail/Yahoo), missing company logo, registration fee mentions, missing location coordinates, and layout anomalies."),
        ("Layout Adaptability", "Identifies if text comes from unstructured formats (PDF/OCR) to avoid penalizing missing field structures.")
    ],
    border_color=COLOR_BLUE
)

# Right Column: Trust bands card
create_card(
    slide_8, Inches(6.8), Inches(1.6), Inches(5.9), Inches(5.0),
    "Trust Scoring Bounds",
    [
        ("90 - 100", "High Trust: Postings from verified recruiters with complete company profiles and zero warnings."),
        ("70 - 89", "Moderate Trust: Minor missing parameters with valid contact info and standard text layout."),
        ("50 - 69", "Low Trust: Multiple warning signals detected. High missingness rates. Job seeker caution advised."),
        ("Below 50", "Unreliable / Highly Suspicious: Severe warning triggers (e.g. fees requested). Extreme missingness.")
    ]
)

add_footer(slide_8, 8)

# =============================================================
# SLIDE 9: Risk & Severity
# =============================================================
slide_9 = prs.slides.add_slide(blank_slide_layout)
apply_theme(slide_9)
add_header(slide_9, "8. Risk & Severity Scorecards")

# Card 1: Severity
create_card(
    slide_9, Inches(0.6), Inches(1.6), Inches(5.8), Inches(5.0),
    "Fraud Severity Score (0-100)",
    [
        ("Impact Metric", "Evaluates the harm a fraudulent listing poses to candidate safety, rather than just raw prediction likelihood."),
        ("6 Severe Signals", "Detects payment checks, bank account setup forms, wire transfer requests, and application processing fee rules."),
        ("Communication Risks", "Deducts major points if recruitment is directed strictly to mobile messaging apps (WhatsApp, Telegram, SMS)."),
        ("Classification Safeguard", "High severity triggers elevate final risk recommendations, even if the ML classifier probability is low.")
    ]
)

# Card 2: Risk categorization matrix
create_card(
    slide_9, Inches(6.8), Inches(1.6), Inches(5.9), Inches(5.0),
    "Risk Categorization Matrix",
    [
        ("Bivariate Mapping", "Maps coordinates of the BERT raw probability and the Heuristic Trust Score to output 4 discrete risk tiers."),
        ("Risk Badges", "Low Risk (Green), Medium Risk (Amber), High Risk (Orange), and Critical Risk (Red)."),
        ("Probability Mitigation", "If BERT classifier output is low (e.g. 0.15) but the trust score is poor, the risk category is escalated to Medium."),
        ("Audit Trails", "Every risk categorization is logged in MongoDB to support developer audit reviews.")
    ],
    border_color=COLOR_BLUE
)

add_footer(slide_9, 9)

# =============================================================
# SLIDE 10: Model Performance Evaluation
# =============================================================
slide_10 = prs.slides.add_slide(blank_slide_layout)
apply_theme(slide_10)
add_header(slide_10, "9. Model Results & Performance")

# Left Column: Performance statistics
create_card(
    slide_10, Inches(0.6), Inches(1.6), Inches(5.8), Inches(5.0),
    "Metrics on Held-Out Test Set (2682 records)",
    [
        ("Accuracy Score", "98.73% classification accuracy on unseen EMSCAD test instances."),
        ("Precision Metric", "90.68% - high precision reduces false positive warnings for legitimate recruiters."),
        ("Recall Metric", "82.31% - leaves room for optimization via focal loss adjustments."),
        ("F1-Score", "0.863 overall F1-score in detecting fraudulent listings."),
        ("ROC-AUC Score", "0.983 area under curve shows strong discrimination performance.")
    ]
)

# Right Column: Confusion Matrix table
tbl_box = slide_10.shapes.add_textbox(Inches(6.8), Inches(1.6), Inches(5.9), Inches(0.5))
tf_tbl = tbl_box.text_frame
p_t = tf_tbl.paragraphs[0]
p_t.text = "Test Set Confusion Matrix"
p_t.font.name = 'Trebuchet MS'
p_t.font.size = Pt(18)
p_t.font.bold = True
p_t.font.color.rgb = COLOR_TEAL

rows, cols = 3, 3
tbl_s = slide_10.shapes.add_table(rows, cols, Inches(6.8), Inches(2.2), Inches(5.933), Inches(2.4))
tbl = tbl_s.table
tbl.columns[0].width = Inches(2.133)
tbl.columns[1].width = Inches(1.9)
tbl.columns[2].width = Inches(1.9)

# Headers
tbl.cell(0, 0).text = ""
tbl.cell(0, 1).text = "Pred Legitimate"
tbl.cell(0, 2).text = "Pred Fraudulent"
for col_idx in range(3):
    cell = tbl.cell(0, col_idx)
    cell.fill.solid()
    cell.fill.fore_color.rgb = COLOR_CARD_BG
    p = cell.text_frame.paragraphs[0]
    p.font.name = 'Trebuchet MS'
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEAL
    p.alignment = PP_ALIGN.CENTER

# Rows
tbl.cell(1, 0).text = "Actual Legitimate"
tbl.cell(2, 0).text = "Actual Fraudulent"
for row_idx in [1, 2]:
    cell = tbl.cell(row_idx, 0)
    cell.fill.solid()
    cell.fill.fore_color.rgb = COLOR_CARD_BG
    p = cell.text_frame.paragraphs[0]
    p.font.name = 'Trebuchet MS'
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_LIGHT
    p.alignment = PP_ALIGN.CENTER

# Matrix content
tbl.cell(1, 1).text = "2,541 (TN)"
tbl.cell(1, 2).text = "11 (FP)"
tbl.cell(2, 1).text = "23 (FN)"
tbl.cell(2, 2).text = "107 (TP)"

for r in [1, 2]:
    for c in [1, 2]:
        cell = tbl.cell(r, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_BG_DARK
        p = cell.text_frame.paragraphs[0]
        p.font.name = 'Calibri'
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEXT_LIGHT
        p.alignment = PP_ALIGN.CENTER

# Matrix comment
comment_box = slide_10.shapes.add_textbox(Inches(6.8), Inches(4.8), Inches(5.933), Inches(1.8))
tf_c = comment_box.text_frame
tf_c.word_wrap = True
p_c = tf_c.paragraphs[0]
p_c.text = "Baseline Validation Summary:"
p_c.font.name = 'Trebuchet MS'
p_c.font.size = Pt(14)
p_c.font.bold = True
p_c.font.color.rgb = COLOR_BLUE
p_c.space_after = Pt(4)

p_c_body = tf_c.add_paragraph()
p_c_body.text = "Accuracy (98.73%) is very close to the base paper baseline (99%). Our recall (82.31%) is slightly lower than baseline (93%) in this run. Precision (90.68%) is highly optimized to avoid warning user candidates unnecessarily."
p_c_body.font.name = 'Calibri'
p_c_body.font.size = Pt(12)
p_c_body.font.color.rgb = COLOR_TEXT_MUTED
p_c_body.line_spacing = 1.25

add_footer(slide_10, 10)

# =============================================================
# SLIDE 11: Key Contributions & Scope
# =============================================================
slide_11 = prs.slides.add_slide(blank_slide_layout)
apply_theme(slide_11)
add_header(slide_11, "10. Key Contribution & Future Scope")

# Left Column Card: Contributions
create_card(
    slide_11, Inches(0.6), Inches(1.6), Inches(5.8), Inches(5.0),
    "Research Novelty",
    [
        ("Dual XAI Integration", "Combines local explainability (LIME) and global explainability (SHAP) with deep BERT classification."),
        ("Bivariate Risk Framework", "Combines neural network probabilities with deterministic heuristics to compute a reliable scorecard."),
        ("Multi-Input Modality", "Supports scraped URLs, parsed PDF files, raw text, and screenshot images (OCR pipelines)."),
        ("Auditable Codebase", "Comprehensive integration tests verify classifier outputs across edits.")
    ],
    border_color=COLOR_BLUE
)

# Right Column Card: Scope
create_card(
    slide_11, Inches(6.8), Inches(1.6), Inches(5.9), Inches(5.0),
    "Future Research Roadmap",
    [
        ("Browser extension", "Real-time in-page scanning tool to analyze listings on LinkedIn and recruitment boards."),
        ("Multilingual Support", "Add tokenizers and classifiers for listings in Spanish, German, French, etc."),
        ("Continual Learning", "Deploy pipelines to automatically update model weights as new scams are reported."),
        ("Platform APIs", "Enable recruitment portals to query listings via validation APIs.")
    ]
)

add_footer(slide_11, 11)

# =============================================================
# SLIDE 12: Conclusion & Q&A
# =============================================================
slide_12 = prs.slides.add_slide(blank_slide_layout)
apply_theme(slide_12)

# Large Title
title_box = slide_12.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.733), Inches(1.5))
tf = title_box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Conclusion & Q&A"
p.font.name = 'Trebuchet MS'
p.font.size = Pt(44)
p.font.bold = True
p.font.color.rgb = COLOR_TEXT_LIGHT

p_sub = tf.add_paragraph()
p_sub.text = "A major step forward in recruitment security, bringing transparency to fake job classification."
p_sub.font.name = 'Calibri'
p_sub.font.size = Pt(18)
p_sub.font.color.rgb = COLOR_BLUE
p_sub.space_before = Pt(8)

# Details Card
create_card(
    slide_12, Inches(0.8), Inches(4.0), Inches(7.5), Inches(2.4),
    "Repository Blueprint Summary",
    [
        "ai_service/     - Python FastAPI model processing logic.",
        "backend/        - Node.js MVC Gateway & MongoDB connection layer.",
        "frontend/       - React Dashboard with interactive LIME visualizations.",
        "docs/ARCHITECTURE.md - Full three-tier architecture definitions."
    ],
    border_color=COLOR_BLUE
)

# Q&A Prompt
qa_box = slide_12.shapes.add_textbox(Inches(8.8), Inches(4.0), Inches(3.7), Inches(2.4))
tf_qa = qa_box.text_frame
tf_qa.word_wrap = True
p_q = tf_qa.paragraphs[0]
p_q.text = "Questions & Discussion"
p_q.font.name = 'Trebuchet MS'
p_q.font.size = Pt(20)
p_q.font.bold = True
p_q.font.color.rgb = COLOR_TEAL
p_q.space_after = Pt(10)

p_q_body = tf_qa.add_paragraph()
p_q_body.text = "We welcome your questions regarding model training, SHAP/LIME integration thresholds, or scoring calculations."
p_q_body.font.name = 'Calibri'
p_q_body.font.size = Pt(13)
p_q_body.font.color.rgb = COLOR_TEXT_MUTED
p_q_body.line_spacing = 1.2

add_footer(slide_12, 12)

# -------------------------------------------------------------
# Save Updated Presentation
# -------------------------------------------------------------
output_path = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "explainable_fake_job_detection_presentation.pptx"))
prs.save(output_path)
print(f"Enhanced presentation generated successfully at:\n{output_path}")
